from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from openpyxl import Workbook

from ard_ossie.adapters.filesystem import RepositoryPaths
from ard_ossie.application.parsing import ParsingService
from ard_ossie.docling_parser import Evidence, ParsedDocument
from ard_ossie.ingestion import SourceFile, SourceRole


class FakeDoclingParser:
    def __init__(self) -> None:
        self.sources: list[SourceFile] = []

    def parse(self, source: SourceFile) -> ParsedDocument:
        self.sources.append(source)
        return ParsedDocument(
            role=source.role,
            source_hash=source.sha256,
            markdown="# Parsed\n",
            evidence=[
                Evidence(
                    source_hash=source.sha256,
                    role=source.role,
                    locator={"document": source.relative_path},
                )
            ],
        )


def test_parse_product_html_returns_hash_parser_kind_and_evidence(tmp_path: Path) -> None:
    """A granular parse result must retain provenance independently of rendered Markdown."""
    source = tmp_path / "products" / "sales" / "sources" / "product-info" / "product.html"
    source.parent.mkdir(parents=True)
    source.write_text("<html>Product</html>", encoding="utf-8")
    parser = FakeDoclingParser()
    service = ParsingService(RepositoryPaths(tmp_path), parser=parser)

    result = service.parse_product_html(source.relative_to(tmp_path))

    assert result.parser_kind == "docling"
    assert result.role == SourceRole.PRODUCT_HTML
    assert result.source_hash == hashlib.sha256(source.read_bytes()).hexdigest()
    assert result.markdown == "# Parsed\n"
    assert result.evidence[0].locator == {
        "document": "products/sales/sources/product-info/product.html"
    }


def test_parse_dictionary_workbook_preserves_cells_and_is_reproducible(tmp_path: Path) -> None:
    """Repeated workbook parsing must produce the same table/column JSON and source hash."""
    source = tmp_path / "dictionary.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data Dictionary"
    sheet.append(
        ["platform", "catalog", "schema", "table", "column", "data_type", "nullable", "pk"]
    )
    sheet.append(["erp", "analytics", "sales", "orders", "order_id", "INT64", False, True])
    workbook.save(source)
    workbook.close()
    service = ParsingService(RepositoryPaths(tmp_path), parser=FakeDoclingParser())

    first = service.parse_dictionary_workbook(Path("dictionary.xlsx"))
    second = service.parse_dictionary_workbook(Path("dictionary.xlsx"))

    assert first.parser_kind == "cell-preserving-excel"
    assert first.tables[0].locator == "erp|analytics|sales|orders"
    assert first.tables[0].columns[0].name == "order_id"
    assert first.model_dump_json() == second.model_dump_json()


def test_source_signature_reads_only_the_required_prefix(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Signature validation must not load an untrusted source file into memory."""
    source = tmp_path / "product.html"
    source.write_bytes(b"<html>Product</html>" + b"x" * 100_000)

    def reject_whole_file_read(path: Path) -> bytes:
        raise AssertionError(f"whole-file read attempted for {path.name}")

    monkeypatch.setattr(Path, "read_bytes", reject_whole_file_read)

    result = ParsingService(
        RepositoryPaths(tmp_path),
        parser=FakeDoclingParser(),
    ).parse_product_html("product.html")

    assert result.role == SourceRole.PRODUCT_HTML
